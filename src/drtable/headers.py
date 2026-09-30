"""Reference headers: rows 2 and 3 of a data sheet show the schema through Excel formulas.

The build never reads rows 2 and 3, so these formulas are only a view for the people editing
data. ``drtable headers`` writes them; it is run on purpose (for new or converted workbooks),
because rewriting a workbook with openpyxl can drop content it does not support.

Excel keeps a link relative only when the schema file sits in the data file's folder or below
it. Otherwise the link stores the absolute path of whoever saved the workbook, and on another
machine the header keeps showing the values from that save. ``link_warnings`` reports that case.
"""

from __future__ import annotations

import os
import posixpath
import re
import zipfile
from pathlib import Path, PurePosixPath
from urllib.parse import unquote

from openpyxl import load_workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.utils import get_column_letter
from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink, ExternalSheetNames

from .errors import ErrorCollector
from .i18n import tr
from .schema import NAME_ROW, SCOPE_ROW, TYPE_ROW
from .schemafile import SCHEMA_SUFFIXES, Schemas
from .sources import find_files, table_name_of

# Parts openpyxl cannot write back; a workbook holding them is left alone.
UNSUPPORTED_PARTS = ("xl/drawings/", "xl/charts/", "xl/pivotTables/", "xl/pivotCache/", "xl/slicers/", "xl/media/")


def write_headers(
    input_path: Path, schema_root: Path, schemas: Schemas, errors: ErrorCollector,
) -> list[Path]:
    """Writes the reference formulas into rows 2-3 of every table sheet under input_path.

    Needs each table's schema file.
    Returns the workbooks changed.
    """
    changed: list[Path] = []
    base = schema_root.parent if schema_root.is_file() else schema_root
    for path, relative in find_files(input_path, allow_empty=True):
        if path.name.lower().endswith(SCHEMA_SUFFIXES):
            continue
        with zipfile.ZipFile(path) as archive:
            blocked = sorted({name.split("/")[1] for name in archive.namelist() if name.startswith(UNSUPPORTED_PARTS)})
        if blocked:
            errors.add(f"[{relative}]", "A1", tr(
                f"그림·차트 등({', '.join(blocked)})이 있어 건드리지 않았습니다. 수식은 다른 시트에서 복사해 넣으세요",
                f"left unchanged: it holds content openpyxl cannot keep ({', '.join(blocked)}). Copy the formulas in by hand",
            ))
            continue
        workbook = load_workbook(path)
        touched = False
        for sheet in workbook.worksheets:
            name = table_name_of(sheet.title)
            schema = schemas.tables.get(name) if name else None
            if schema is None:
                continue
            view = schema
            target = Path(os.path.relpath(base / view.file, path.parent)).as_posix()
            index = _link_index(workbook, target, view.title or view.name)
            reference = f"'[{index}]{view.title or view.name}'"
            missing = tr("(스키마에 없음)", "(not in schema)")
            for column in range(1, sheet.max_column + 1):
                field_name = sheet.cell(NAME_ROW, column).value
                if field_name in (None, ""):
                    break
                if str(field_name).strip().startswith("#"):
                    continue
                letter = get_column_letter(column)
                for row, source in ((TYPE_ROW, "B"), (SCOPE_ROW, "C")):
                    sheet.cell(row, column).value = (
                        f"=IFERROR(INDEX({reference}!${source}:${source},"
                        f"MATCH({letter}${NAME_ROW},{reference}!$A:$A,0)),\"{missing}\")"
                    )
            touched = True
        if touched:
            workbook.save(path)
            changed.append(path)
    return changed


def link_warnings(path: Path, relative: str, schema_files: set[Path]) -> list[str]:
    """Warns when a data workbook links to a schema file that is not at the linked path here."""
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            messages = []
            for name in names:
                if not re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", name):
                    continue
                book = re.search(rb'<externalBook[^>]*\br:id="([^"]+)"', archive.read(name))
                rels_name = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels")
                if book is None or rels_name not in names:
                    continue
                rels = archive.read(rels_name).decode("utf-8", "replace")
                found = re.search(rf'<Relationship[^>]*Id="{re.escape(book.group(1).decode())}"[^>]*/>', rels)
                target = re.search(r'Target="([^"]+)"', found.group(0)).group(1) if found else ""
                if not target.lower().endswith(SCHEMA_SUFFIXES):
                    continue
                resolved = _resolve(target, path.parent)
                if resolved is None or resolved not in schema_files:
                    messages.append(tr(
                        f"[{relative}]: 참고 헤더가 이 PC에 없는 경로를 가리킵니다({unquote(target)}). "
                        "헤더가 스키마 변경을 따라가지 않으니 drtable headers로 다시 연결하세요",
                        f"[{relative}]: the reference headers link to a path that does not exist here "
                        f"({unquote(target)}), so they do not follow schema changes; relink with drtable headers",
                    ))
            return messages
    except (OSError, zipfile.BadZipFile, KeyError):
        return []


def _resolve(target: str, folder: Path) -> Path | None:
    text = unquote(target)
    if text.lower().startswith("file:///"):
        text = text[len("file:///"):]
    candidate = Path(text.replace("\\", "/"))
    if not re.match(r"^[A-Za-z]:", text) and not text.startswith(("/", "\\")):
        candidate = folder / PurePosixPath(text.replace("\\", "/"))
    try:
        return candidate.resolve() if candidate.exists() else None
    except OSError:
        return None


def _link_index(workbook: object, target: str, sheet_name: str) -> int:
    """1-based index of the workbook's external link to target, adding the link if needed."""
    links = workbook._external_links
    for position, link in enumerate(links, start=1):
        if link.file_link is not None and link.file_link.Target == target:
            return position
    link = ExternalLink(externalBook=ExternalBook(sheetNames=ExternalSheetNames(sheetName=[sheet_name]), id="rId1"))
    link.file_link = Relationship(type="externalLinkPath", Target=target, TargetMode="External", Id="rId1")
    links.append(link)
    return len(links)
