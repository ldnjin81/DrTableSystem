"""Schema files: where the fields of a table or an enum are defined.

A table's schema is either ``<Name>.schema.xlsx`` or ``<Name>.schema.yaml``. When both exist
the YAML file is the source and the xlsx file is a view generated from it (``drtable
schema-export``) so that data workbooks can show the fields through formulas.

xlsx layout: one definition sheet named like the data sheet (``Items`` or ``<enum>Kind``);
row 1 holds labels, and from row 2 each row is a field: A name, B type, C scope, D comment.

YAML layout::

    table: Items            # or  enum: Kind
    fields:
      - name: Id
        type: ID<int32>
        scope: all
      - name: Cost
        type: int32=0
        scope: server
        comment: purchase price
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .errors import ErrorCollector
from .i18n import tr
from .schema import IDENTIFIER_RE
from .sources import ENUM_SHEET_RE, find_files, strip_sheet_comment

XLSX_SUFFIX = ".schema.xlsx"
YAML_SUFFIXES = (".schema.yaml", ".schema.yml")
SCHEMA_SUFFIXES = (XLSX_SUFFIX, *YAML_SUFFIXES)
LABELS = ("Field", "Type", "Scope", "Comment")

# (row or line, name, type, scope, comment)
Field = tuple[int, object, object, object, object]


@dataclass
class Schema:
    """The field definitions of one table or enum, read from a schema file."""

    file: str
    title: str  # "Items" or "<enum>Kind" (the xlsx sheet name, comments included)
    fields: list[Field]
    kind: str = "xlsx"  # "xlsx" or "yaml"
    view: Schema | None = field(default=None, repr=False)  # the xlsx view of a YAML schema
    # YAML: (field line, header row 1/2/3) -> line of its name/type/scope key.
    lines: dict[tuple[int, int], int] = field(default_factory=dict, repr=False)

    @property
    def key(self) -> str:
        return strip_sheet_comment(self.title)

    @property
    def name(self) -> str:
        match = ENUM_SHEET_RE.fullmatch(self.key)
        return match.group("name") if match else self.key

    @property
    def where(self) -> str:
        return self.file if self.kind == "yaml" else f"[{self.file}]{self.title}"

    def cell(self, row: int, header_row: int) -> str:
        """Cell (xlsx: A name, B type, C scope) or line (YAML) of a field definition."""
        if self.kind == "yaml":
            return str(self.lines.get((row, header_row), row))
        return f"{'ABC'[header_row - 1]}{row}"

    def at(self, row: int, header_row: int = 1) -> str:
        separator = ":" if self.kind == "yaml" else "!"
        return f"{self.where}{separator}{self.cell(row, header_row)}"

    def raw_columns(self) -> list[tuple[int, object, object, object]]:
        return [(row, name, type_, scope) for row, name, type_, scope, _ in self.fields]

    def definition(self) -> list[tuple[str, str, str, str]]:
        """Field values only, for comparing a YAML source with its xlsx view."""
        return [
            tuple("" if value is None else str(value).strip() for value in item[1:])
            for item in self.fields
        ]


def load_schemas(root: Path, errors: ErrorCollector) -> dict[str, Schema]:
    """Every schema file under root, keyed by 'Items' or '<enum>Kind'.

    A YAML file wins over an xlsx file of the same name; that xlsx file becomes its view.
    """
    schemas: dict[str, Schema] = {}
    views: dict[str, Schema] = {}
    for path, relative in find_files(root, SCHEMA_SUFFIXES, allow_empty=True):
        is_yaml = path.name.lower().endswith(YAML_SUFFIXES)
        schema = _read_yaml(path, relative, errors) if is_yaml else _read_xlsx(path, relative, errors)
        if schema is None:
            continue
        stem = path.name[: -len(next(s for s in SCHEMA_SUFFIXES if path.name.lower().endswith(s)))]
        if schema.name != stem:
            errors.add(schema.where, schema.cell(1, 1) if is_yaml else "A1", tr(
                f"이름 '{schema.name}'과 파일 이름 '{path.name}'이 다릅니다. 파일 이름은 '{schema.name}{path.name[len(stem):]}'여야 합니다",
                f"name '{schema.name}' does not match the file name '{path.name}'; name the file '{schema.name}{path.name[len(stem):]}'",
            ))
            continue
        bucket = schemas if is_yaml else views
        if schema.key in bucket:
            errors.add(schema.where, "1" if is_yaml else "A1", tr(
                f"'{schema.key}'의 스키마가 {bucket[schema.key].where}에도 있습니다",
                f"'{schema.key}' also has a schema in {bucket[schema.key].where}",
            ))
            continue
        bucket[schema.key] = schema
    for key, view in views.items():
        if key in schemas:
            schemas[key].view = view
        else:
            schemas[key] = view
    return schemas


def stale_views(schemas: dict[str, Schema]) -> list[str]:
    """Warnings for xlsx views that no longer match their YAML source."""
    return [
        tr(
            f"{schema.view.where}: {schema.file}와 내용이 다릅니다. drtable schema-export로 다시 만드세요",
            f"{schema.view.where}: does not match {schema.file}; regenerate it with drtable schema-export",
        )
        for schema in schemas.values()
        if schema.kind == "yaml" and schema.view is not None
        and schema.view.definition() != schema.definition()
    ]


def export_views(root: Path, check: bool = False) -> tuple[list[Path], list[str]]:
    """Writes <Name>.schema.xlsx next to every <Name>.schema.yaml under root.

    Output is deterministic, and unchanged files are not rewritten. With ``check`` nothing is
    written; the files that would change are returned instead (for CI).
    Returns (changed files, errors).
    """
    errors = ErrorCollector()
    changed: list[Path] = []
    for path, relative in find_files(root, YAML_SUFFIXES, allow_empty=True):
        schema = _read_yaml(path, relative, errors)
        if schema is None:
            continue
        stem = path.name[: -len(next(s for s in YAML_SUFFIXES if path.name.lower().endswith(s)))]
        target = path.with_name(f"{stem}{XLSX_SUFFIX}")
        data = render_xlsx(schema.title, schema.fields)
        if target.exists() and target.read_bytes() == data:
            continue
        changed.append(target)
        if not check:
            target.write_bytes(data)
    return changed, errors.messages


def render_xlsx(title: str, fields: list[Field]) -> bytes:
    """A schema workbook as bytes. The same fields always give the same bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(list(LABELS))
    for _, name, type_, scope, comment in fields:
        sheet.append([name, type_, scope, comment])
    for column, width in zip("ABCD", (24, 28, 10, 40), strict=True):
        sheet.column_dimensions[column].width = width
    fixed = datetime(2000, 1, 1)  # noqa: DTZ001 - openpyxl writes naive UTC times
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.properties.lastModifiedBy = None
    workbook.properties.creator = "DrTableSystem"
    buffer = io.BytesIO()
    workbook.save(buffer)
    # Zip entries carry the time they were written; pin it so the bytes are reproducible.
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            pinned = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            pinned.compress_type = zipfile.ZIP_DEFLATED
            pinned.external_attr = info.external_attr
            target.writestr(pinned, source.read(info.filename))
    return output.getvalue()


def dump_yaml(title: str, fields: list[Field]) -> str:
    """A schema as YAML text (used when converting to the YAML source layout)."""
    match = ENUM_SHEET_RE.fullmatch(strip_sheet_comment(title))
    kind, name = ("enum", match.group("name")) if match else ("table", strip_sheet_comment(title))
    lines = [f"{kind}: {name}", "fields:"]
    for _, field_name, type_, scope, comment in fields:
        lines.append(f"  - name: {_yaml_scalar(field_name)}")
        lines.append(f"    type: {_yaml_scalar(type_)}")
        lines.append(f"    scope: {_yaml_scalar(scope)}")
        if comment not in (None, ""):
            lines.append(f"    comment: {_yaml_scalar(comment)}")
    return "\n".join(lines) + "\n"


def _yaml_scalar(value: object) -> str:
    text = "" if value is None else str(value)
    plain = yaml.safe_dump(text, allow_unicode=True, default_flow_style=True, width=1 << 16)
    plain = plain.removesuffix("\n...\n").removesuffix("\n")
    return plain


def _read_xlsx(path: Path, relative: str, errors: ErrorCollector) -> Schema | None:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
        errors.add(f"[{relative}]", "A1", tr(f"xlsx 파일을 읽을 수 없습니다: {exc}", f"cannot read the xlsx file: {exc}"))
        return None
    try:
        sheets = [sheet for sheet in workbook.worksheets if not sheet.title.startswith("#")]
        if len(sheets) != 1:
            errors.add(f"[{relative}]", "A1", tr(
                "스키마 파일에는 정의 시트가 하나만 있어야 합니다('#' 메모 시트 제외)",
                "a schema file needs exactly one definition sheet ('#' note sheets aside)",
            ))
            return None
        sheet = sheets[0]
        schema = Schema(relative, sheet.title, [])
        if not _valid_name(schema, errors):
            return None
        for row, values in enumerate(sheet.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
            values = tuple(values) + (None,) * (4 - len(values))
            if all(value in (None, "") for value in values[:3]):
                continue
            if values[0] in (None, ""):
                errors.add(schema.where, f"A{row}", tr("필드명이 비어 있습니다", "the field name is empty"))
                continue
            schema.fields.append((row, *values))
        return schema
    finally:
        workbook.close()


def _read_yaml(path: Path, relative: str, errors: ErrorCollector) -> Schema | None:
    try:
        text = path.read_text(encoding="utf-8")
        root = yaml.compose(text)
    except (OSError, UnicodeDecodeError) as exc:
        errors.add(relative, "1", tr(f"파일을 읽을 수 없습니다: {exc}", f"cannot read the file: {exc}"))
        return None
    except yaml.YAMLError as exc:
        line = getattr(getattr(exc, "problem_mark", None), "line", 0) + 1
        errors.add(relative, str(line), tr(f"YAML 문법 오류: {exc}", f"YAML syntax error: {exc}"))
        return None
    if not isinstance(root, yaml.MappingNode):
        errors.add(relative, "1", tr("'table:' 또는 'enum:'과 'fields:'가 있는 매핑이어야 합니다", "expected a mapping with 'table:' or 'enum:' and 'fields:'"))
        return None
    top = {key.value: value for key, value in root.value}
    kinds = [kind for kind in ("table", "enum") if kind in top]
    if len(kinds) != 1 or not isinstance(top[kinds[0]], yaml.ScalarNode):
        errors.add(relative, "1", tr("'table: 이름'과 'enum: 이름' 중 하나가 필요합니다", "exactly one of 'table: Name' or 'enum: Name' is required"))
        return None
    name = str(top[kinds[0]].value).strip()
    title = f"<enum>{name}" if kinds[0] == "enum" else name
    schema = Schema(relative, title, [], kind="yaml")
    if not _valid_name(schema, errors):
        return None
    for key in top.keys() - {"table", "enum", "fields"}:
        errors.add(relative, "1", tr(f"알 수 없는 키 '{key}'", f"unknown key '{key}'"))
    fields = top.get("fields")
    if not isinstance(fields, yaml.SequenceNode):
        errors.add(relative, "1", tr("'fields:' 목록이 필요합니다", "a 'fields:' list is required"))
        return None
    for item in fields.value:
        line = item.start_mark.line + 1
        if not isinstance(item, yaml.MappingNode):
            errors.add(relative, str(line), tr("필드는 name/type/scope를 가진 매핑이어야 합니다", "a field must be a mapping with name/type/scope"))
            continue
        values: dict[str, object] = {}
        key_rows = {"name": 1, "type": 2, "scope": 3}
        for key, value in item.value:
            if key.value in key_rows:
                schema.lines[(line, key_rows[key.value])] = value.start_mark.line + 1
            if key.value not in ("name", "type", "scope", "comment"):
                errors.add(relative, str(key.start_mark.line + 1), tr(f"알 수 없는 필드 키 '{key.value}'", f"unknown field key '{key.value}'"))
            elif not isinstance(value, yaml.ScalarNode):
                errors.add(relative, str(value.start_mark.line + 1), tr(f"'{key.value}'에는 값 하나만 씁니다", f"'{key.value}' takes a single value"))
            else:
                values[key.value] = None if value.tag.endswith(":null") else value.value
        if values.get("name") in (None, ""):
            errors.add(relative, str(line), tr("필드명이 비어 있습니다", "the field name is empty"))
            continue
        if values.get("scope") is None:
            # '#' starts a YAML comment, so an unquoted  scope: #  reads as empty.
            errors.add(relative, schema.cell(line, 3), tr(
                "scope가 비어 있습니다. '#'은 YAML 주석 기호라 scope: \"#\"처럼 따옴표로 감싸야 합니다",
                "scope is empty. '#' starts a YAML comment, so write it quoted: scope: \"#\"",
            ))
            continue
        schema.fields.append((line, values["name"], values.get("type"), values["scope"], values.get("comment")))
    return schema


def _valid_name(schema: Schema, errors: ErrorCollector) -> bool:
    if IDENTIFIER_RE.fullmatch(schema.name) and (schema.key == schema.name or ENUM_SHEET_RE.fullmatch(schema.key)):
        return True
    errors.add(schema.where, "1" if schema.kind == "yaml" else "A1", tr(
        f"올바르지 않은 이름 '{schema.title}'", f"invalid name '{schema.title}'",
    ))
    return False
