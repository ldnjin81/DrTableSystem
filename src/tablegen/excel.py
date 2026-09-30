"""Reads .xlsx files into a validated data model."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .errors import ErrorCollector, ValidationErrors
from .i18n import tr
from .schema import (
    DATA_ROW,
    IDENTIFIER_RE,
    NAME_ROW,
    SCOPE_ROW,
    TYPE_ROW,
    ColumnSchema,
    EnumSchema,
    EnumValue,
    TableSchema,
    TableSource,
    build_columns,
    calculate_schema_hash,
    parse_type,
)
from .values import convert_value

ENUM_SHEET_RE = re.compile(r"^<enum>(?P<name>[A-Za-z][A-Za-z0-9_]*)$")


@dataclass(frozen=True)
class DataModel:
    source_files: tuple[str, ...]
    enums: tuple[EnumSchema, ...]
    tables: tuple[TableSchema, ...]


@dataclass(frozen=True)
class SheetRef:
    """A worksheet together with the file it came from (path relative to the input folder)."""

    file: str
    sheet: object

    @property
    def where(self) -> str:
        """Excel-style location used in every message: [File.xlsx]Sheet."""
        return f"[{self.file}]{self.sheet.title}"


@dataclass
class _RowState:
    """Shared across the parts of one table so duplicates are found between files too."""

    used_keys: dict[object, str] = field(default_factory=dict)
    name_spellings: dict[tuple[str, str], tuple[str, str]] = field(default_factory=dict)


# '#' starts a comment everywhere: a sheet named "#Notes" is skipped, and in "Items#Weapons"
# the text from '#' on is a comment, so the sheet is part of table Items.
SHEET_COMMENT = "#"


def strip_sheet_comment(sheet_title: str) -> str:
    """The sheet name without its #comment and surrounding spaces."""
    return sheet_title.split(SHEET_COMMENT, 1)[0].strip()


def table_name_of(sheet_title: str) -> str | None:
    """Table name for a table sheet, or None for notes and enums.

    The sheet name without its #comment is the table name. Sheets with the same table name,
    in one file or in several, are parts of one table.
    """
    if sheet_title.startswith("#") or sheet_title.startswith("<enum>"):
        return None
    return strip_sheet_comment(sheet_title)


def load_model(input_path: Path) -> DataModel:
    files = _find_files(input_path)
    errors = ErrorCollector()
    workbooks: list[tuple[str, object]] = []
    for path, relative in files:
        try:
            workbooks.append((relative, load_workbook(path, data_only=True, read_only=True)))
        except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
            errors.add(f"[{relative}]", "A1", tr(f"xlsx 파일을 읽을 수 없습니다: {exc}", f"cannot read the xlsx file: {exc}"))
    errors.raise_if_any()

    enum_sheets: dict[str, SheetRef] = {}
    table_parts: dict[str, list[SheetRef]] = {}
    for relative, workbook in workbooks:
        for sheet in workbook.worksheets:
            ref = SheetRef(relative, sheet)
            if sheet.title.startswith("#"):
                continue
            if sheet.title.startswith("<enum>"):
                match = ENUM_SHEET_RE.fullmatch(strip_sheet_comment(sheet.title))
                if not match:
                    errors.add(ref.where, "A1", tr(
                        f"올바르지 않은 열거형 시트 이름 '{sheet.title}'",
                        f"invalid enum sheet name '{sheet.title}'",
                    ))
                    continue
                name = match.group("name")
                if name in enum_sheets:
                    errors.add(ref.where, "A1", tr(
                        f"열거형 '{name}'이 {enum_sheets[name].where}에도 정의되어 있습니다 (열거형은 나눌 수 없습니다)",
                        f"enum '{name}' is also defined in {enum_sheets[name].where} (enums cannot be split)",
                    ))
                else:
                    enum_sheets[name] = ref
                continue
            name = table_name_of(sheet.title)
            if not name or not IDENTIFIER_RE.fullmatch(name):
                errors.add(ref.where, "A1", tr(
                    f"올바르지 않은 테이블 이름 '{sheet.title}'",
                    f"invalid table name '{sheet.title}'",
                ))
                continue
            table_parts.setdefault(name, []).append(ref)

    # Register every enum name first so that enum sheets can reference each other.
    enum_names = {
        name: EnumSchema(name, ref.sheet.title, ref.file, ())
        for name, ref in enum_sheets.items()
    }
    enums: dict[str, EnumSchema] = {}
    enum_columns: dict[str, list[ColumnSchema]] = {}
    for name, ref in enum_sheets.items():
        enum, columns = _parse_enum(ref, name, enum_names, errors)
        if enum:
            enums[name] = enum
            enum_columns[name] = columns

    # Resolve every table's primary key type before reading rows (Ref<T> needs them).
    # The first part of a table defines its header; the other parts are checked against it.
    table_keys: dict[str, tuple[str, str, str]] = {}
    for name, parts in table_parts.items():
        first = parts[0]
        for index, raw_name, raw_type, raw_scope in _raw_columns(first.sheet):
            scope = str(raw_scope).strip().lower()
            if scope == "#":
                continue
            if not re.match(r"^(id|subkey)\s*<", str(raw_type).strip(), re.IGNORECASE):
                if IDENTIFIER_RE.fullmatch(str(raw_name)):
                    table_keys[f"{name}.{raw_name}"] = (str(raw_type), "field", scope)
                continue
            parsed = parse_type(raw_type, first.where, f"{get_column_letter(index)}{TYPE_ROW}", errors)
            if parsed and parsed.role == "id":
                table_keys[name] = (parsed.type_name, "id", scope)
                table_keys[f"{name}.{raw_name}"] = (parsed.type_name, "id", scope)
            elif parsed and parsed.role == "subkey":
                table_keys[f"{name}.{raw_name}"] = (parsed.type_name, "subkey", scope)
    for name, columns in enum_columns.items():
        if any(column.role != "id" and column.name != "Value" for column in columns):
            info_name = f"{name}Info"
            table_keys[info_name] = (f"E{name}", "id", "all")
            for column in columns:
                if column.role == "id":
                    table_keys[f"{info_name}.{column.name}"] = (f"E{name}", "id", "all")
                elif column.name != "Value":
                    table_keys[f"{info_name}.{column.name}"] = (
                        column.type_name, column.role or "field", column.scope
                    )

    tables: list[TableSchema] = []
    for name, parts in table_parts.items():
        table = _parse_table(name, parts, enums, table_keys, errors)
        if table:
            tables.append(table)

    table_names = set(table_parts)
    for name, ref in enum_sheets.items():
        enum = enums.get(name)
        if enum is None or name not in enum_columns:
            continue
        columns = build_columns(_raw_columns(ref.sheet), ref.where, enums, errors, table_keys)
        info_columns = [column for column in columns if column.name != "Value"]
        extras = [column for column in info_columns if column.role != "id"]
        if not extras:
            continue
        info_name = f"{name}Info"
        if info_name in table_names:
            errors.add(ref.where, "A1", tr(f"Info 테이블 이름 '{info_name}'이 중복되었습니다", f"Info table name '{info_name}' is already used"))
            continue
        table_names.add(info_name)
        transformed = [
            ColumnSchema(
                column.name,
                f"E{name}" if column.role == "id" else column.type_name,
                column.role,
                column.scope,
                column.source_columns,
                column.header_cells,
                column.array_size,
                column.default_values,
                column.ref_target,
                column.ref_key,
            )
            for column in info_columns
        ]
        info_table = TableSchema(name=info_name, sheet=ref.sheet.title, source_name=ref.file, columns=transformed)
        rows = _read_rows(info_table, ref, transformed, enums, errors, _RowState())
        info_table.sources.append(TableSource(ref.file, ref.sheet.title, rows))
        info_table.schema_hash = calculate_schema_hash(transformed, enums)
        tables.append(info_table)

    for _, workbook in workbooks:
        workbook.close()
    errors.raise_if_any()
    return DataModel(
        tuple(relative for relative, _ in workbooks),
        tuple(sorted(enums.values(), key=lambda item: item.name)),
        tuple(sorted(tables, key=lambda item: item.name)),
    )


def _find_files(input_path: Path) -> list[tuple[Path, str]]:
    """The .xlsx files to read, with their paths relative to the input folder (sorted)."""
    if input_path.is_file() and input_path.suffix.lower() == ".xlsx":
        return [(input_path, input_path.name)]
    if input_path.is_dir():
        found = []
        for path in input_path.rglob("*"):
            relative = path.relative_to(input_path)
            if (
                path.is_file()
                and path.suffix.lower() == ".xlsx"
                # Skip Excel lock files (~$Book.xlsx) and hidden folders such as .git.
                and not path.name.startswith("~$")
                and not any(part.startswith(".") for part in relative.parts)
            ):
                found.append((path, relative.as_posix()))
        if found:
            return sorted(found, key=lambda item: item[1])
    raise ValidationErrors([tr(f"입력!A1: xlsx 파일을 찾을 수 없습니다: {input_path}", f"input!A1: no xlsx file found: {input_path}")])


def _parse_enum(
    ref: SheetRef,
    name: str,
    enum_names: dict[str, EnumSchema],
    errors: ErrorCollector,
) -> tuple[EnumSchema | None, list[ColumnSchema]]:
    sheet = ref.sheet
    where = ref.where
    raw_columns = _raw_columns(sheet)
    columns = build_columns(raw_columns, where, enum_names, errors)
    ids = [column for column in columns if column.role == "id"]
    if len(ids) != 1:
        return None, columns
    primary = ids[0]
    if primary.type_name != "name" or primary.is_array:
        errors.add(where, _cell_for(primary, TYPE_ROW), tr("열거형 기본키는 ID<name>이어야 합니다", "an enum sheet key must be ID<name>"))
    value_columns = [column for column in columns if column.name == "Value"]
    value_column = value_columns[0] if value_columns else None
    if value_column and (
        value_column.type_name != "int32" or value_column.role is not None or value_column.is_array
    ):
        errors.add(where, _cell_for(value_column, TYPE_ROW), tr("Value 열은 int32 일반 필드여야 합니다", "the Value column must be a plain int32 field"))
    comment_column = next(
        (
            column_index
            for column_index, raw_name, _, raw_scope in raw_columns
            if str(raw_name).strip() == "Comment"
            and str(raw_scope).strip().upper() == "#"
        ),
        None,
    )
    values: list[EnumValue] = []
    used_names: set[str] = set()
    used_values: set[int] = set()
    next_value = 0
    active_columns = {source for column in columns for source in column.source_columns}
    if comment_column:
        active_columns.add(comment_column)
    for row in range(DATA_ROW, sheet.max_row + 1):
        if all(sheet.cell(row, column).value in (None, "") for column in active_columns):
            continue
        raw_name = sheet.cell(row, primary.source_columns[0]).value
        raw_value = (
            sheet.cell(row, value_column.source_columns[0]).value if value_column else None
        )
        raw_comment = sheet.cell(row, comment_column).value if comment_column else None
        item_name = "" if raw_name is None else str(raw_name).strip()
        if not IDENTIFIER_RE.fullmatch(item_name):
            errors.add(
                where,
                f"{get_column_letter(primary.source_columns[0])}{row}",
                tr(f"올바르지 않은 열거자 이름 '{item_name}'", f"invalid enumerator name '{item_name}'"),
            )
            continue
        if item_name in used_names:
            errors.add(
                where,
                f"{get_column_letter(primary.source_columns[0])}{row}",
                tr(f"열거자 이름 '{item_name}'이 중복되었습니다", f"enumerator '{item_name}' appears twice"),
            )
        used_names.add(item_name)
        if raw_value in (None, ""):
            if value_column and value_column.default_values[0] is not None:
                value = int(value_column.default_values[0])
            else:
                value = next_value
        else:
            try:
                value = int(raw_value)
                if isinstance(raw_value, float) and not raw_value.is_integer():
                    raise ValueError
            except (TypeError, ValueError):
                cell = _cell_for(value_column, row) if value_column else "A1"
                errors.add(where, cell, tr(f"열거형 값 '{raw_value}'은 정수가 아닙니다", f"enum value '{raw_value}' is not an integer"))
                continue
        if not 0 <= value <= 255:
            cell = _cell_for(value_column, row) if value_column else _cell_for(primary, row)
            errors.add(where, cell, tr("열거형 값은 uint8 범위(0~255)여야 합니다", "enum values must fit in uint8 (0-255)"))
        if value in used_values:
            cell = _cell_for(value_column, row) if value_column else _cell_for(primary, row)
            errors.add(where, cell, tr(f"열거형 값 {value}가 중복되었습니다", f"enum value {value} appears twice"))
        used_values.add(value)
        next_value = value + 1
        values.append(EnumValue(item_name, value, "" if raw_comment is None else str(raw_comment)))
    if not values:
        errors.add(where, "A4", tr("열거형에는 항목이 하나 이상 필요합니다", "an enum needs at least one value"))
        return None, columns
    return EnumSchema(name, sheet.title, ref.file, tuple(values)), columns


def _parse_table(
    name: str,
    parts: list[SheetRef],
    enums: dict[str, EnumSchema],
    table_keys: dict[str, tuple[str, str, str]],
    errors: ErrorCollector,
) -> TableSchema | None:
    first = parts[0]
    columns = build_columns(_raw_columns(first.sheet), first.where, enums, errors, table_keys)
    if not columns or not any(column.role == "id" for column in columns):
        return None
    table = TableSchema(name=name, sheet=first.sheet.title, source_name=first.file, columns=columns)
    state = _RowState()
    expected = {column.name: _field_signature(column) for column in columns}
    for part in parts:
        part_columns = columns
        if part is not first:
            part_columns = build_columns(_raw_columns(part.sheet), part.where, enums, errors, table_keys)
            if not _same_fields(part, part_columns, first, expected, errors):
                continue
        rows = _read_rows(table, part, part_columns, enums, errors, state)
        table.sources.append(TableSource(part.file, part.sheet.title, rows))
    table.schema_hash = calculate_schema_hash(columns, enums)
    return table


def _field_signature(column: ColumnSchema) -> tuple[object, ...]:
    return (
        column.type_name, column.role, column.scope, column.array_size,
        column.default_values, column.ref_target, column.ref_key,
    )


def _same_fields(
    part: SheetRef,
    columns: list[ColumnSchema],
    first: SheetRef,
    expected: dict[str, tuple[object, ...]],
    errors: ErrorCollector,
) -> bool:
    """Every part of a split table must declare the same fields; column order may differ."""
    actual = {column.name: column for column in columns}
    ok = True
    for field_name in expected.keys() - actual.keys():
        errors.add(part.where, "A1", tr(
            f"필드 '{field_name}'이 없습니다. 같은 테이블의 시트는 모두 {first.where}와 같은 필드를 가져야 합니다",
            f"field '{field_name}' is missing. Every sheet of a split table needs the same fields as {first.where}",
        ))
        ok = False
    for field_name, column in actual.items():
        if field_name not in expected:
            errors.add(part.where, column.header_cells[0], tr(
                f"필드 '{field_name}'이 {first.where}에 없습니다. 같은 테이블의 시트는 모두 같은 필드를 가져야 합니다",
                f"field '{field_name}' is not in {first.where}. Every sheet of a split table needs the same fields",
            ))
            ok = False
        elif _field_signature(column) != expected[field_name]:
            errors.add(part.where, column.header_cells[0], tr(
                f"필드 '{field_name}'의 타입·범위·배열 크기·기본값이 {first.where}와 다릅니다",
                f"field '{field_name}' differs from {first.where} in type, scope, array size or default",
            ))
            ok = False
    return ok


def _read_rows(
    table: TableSchema,
    ref: SheetRef,
    columns: list[ColumnSchema],
    enums: dict[str, EnumSchema],
    errors: ErrorCollector,
    state: _RowState,
) -> int:
    """Appends the rows of one sheet to the table and returns how many were read."""
    sheet = ref.sheet
    where = ref.where
    primary = next(column for column in columns if column.role == "id")
    sub_keys = [column for column in columns if column.role == "subkey"]
    used_keys = state.used_keys
    # (field, lowercase value) -> (first spelling, location). Unreal FName is case-insensitive,
    # so 'Sword' and 'sword' would collapse into one key and break the sort contract.
    name_spellings = state.name_spellings
    data_source_columns = {
        source_column for column in columns for source_column in column.source_columns
    }
    count = 0
    for row in range(DATA_ROW, sheet.max_row + 1):
        if all(sheet.cell(row, index).value in (None, "") for index in data_source_columns):
            continue
        converted: dict[str, object] = {}
        for column in columns:
            if column.is_array:
                for source_column in column.source_columns:
                    raw_value = sheet.cell(row, source_column).value
                    if column.ref_target and column.type_name.startswith("E") and raw_value in (None, ""):
                        errors.add(
                            where,
                            f"{get_column_letter(source_column)}{row}",
                            tr("열거형 기본키를 참조하는 셀은 비울 수 없습니다", "a reference to an enum-keyed table cannot be empty"),
                        )
                converted[column.name] = [
                    convert_value(
                        _value_with_default(
                            sheet.cell(row, source_column).value,
                            column.default_values[position],
                        ),
                        column.type_name,
                        enums,
                        where,
                        f"{get_column_letter(source_column)}{row}",
                        errors,
                    )
                    for position, source_column in enumerate(column.source_columns)
                ]
            else:
                source_column = column.source_columns[0]
                raw_value = sheet.cell(row, source_column).value
                if column.ref_target and column.type_name.startswith("E") and raw_value in (None, ""):
                    errors.add(
                        where,
                        f"{get_column_letter(source_column)}{row}",
                        tr("열거형 기본키를 참조하는 셀은 비울 수 없습니다", "a reference to an enum-keyed table cannot be empty"),
                    )
                if column.role == "id" and raw_value in (None, ""):
                    errors.add(
                        where,
                        f"{get_column_letter(source_column)}{row}",
                        tr("기본키 값이 비어 있습니다", "the primary key is empty"),
                    )
                converted[column.name] = convert_value(
                    _value_with_default(raw_value, column.default_values[0]),
                    column.type_name,
                    enums,
                    where,
                    f"{get_column_letter(source_column)}{row}",
                    errors,
                )
        key = converted[primary.name]
        key_cell = f"{get_column_letter(primary.source_columns[0])}{row}"
        if key in used_keys:
            seen = used_keys[key]
            errors.add(where, key_cell, tr(
                f"기본키 값 '{key}'이 중복되었습니다" + ("" if seen.startswith(where + "!") else f" (처음: {seen})"),
                f"duplicate primary key '{key}'" + ("" if seen.startswith(where + "!") else f" (first at {seen})"),
            ))
        else:
            used_keys[key] = f"{where}!{key_cell}"
            if primary.type_name == "name":
                _check_name_case(where, primary.name, key, key_cell, name_spellings, errors)
        for column in sub_keys:
            if column.type_name == "name" and column.name != primary.name:
                cell = f"{get_column_letter(column.source_columns[0])}{row}"
                _check_name_case(where, column.name, converted[column.name], cell, name_spellings, errors)
        # Keep the field order of the first part so every row looks the same.
        table.rows.append({column.name: converted[column.name] for column in table.columns})
        count += 1
    return count


def _check_name_case(
    sheet: str,
    field: str,
    value: object,
    cell: str,
    spellings: dict[tuple[str, str], tuple[str, str]],
    errors: ErrorCollector,
) -> None:
    """Reports name key values that differ from an earlier value only by case."""
    if not isinstance(value, str) or value == "":
        return
    folded = (field, value.lower())
    seen = spellings.get(folded)
    if seen is None:
        spellings[folded] = (value, f"{sheet}!{cell}")
    elif seen[0] != value:
        errors.add(
            sheet,
            cell,
            tr(
                f"name 키 '{value}'이 {seen[1]}의 '{seen[0]}'과 대소문자만 다릅니다. "
                "언리얼 FName은 대소문자를 구분하지 않아 같은 값이 됩니다",
                f"name key '{value}' differs from '{seen[0]}' at {seen[1]} only by case. "
                "Unreal FName is case-insensitive, so they would be the same value",
            ),
        )


def _value_with_default(value: object, declared_default: object | None) -> object:
    if value in (None, "") and declared_default is not None:
        return declared_default
    return value


def _raw_columns(sheet: object) -> list[tuple[int, object, object, object]]:
    columns: list[tuple[int, object, object, object]] = []
    for column in range(1, sheet.max_column + 1):
        name = sheet.cell(NAME_ROW, column).value
        if name in (None, ""):
            break
        columns.append(
            (column, name, sheet.cell(TYPE_ROW, column).value, sheet.cell(SCOPE_ROW, column).value)
        )
    return columns


def _cell_for(column: ColumnSchema | None, row: int) -> str:
    if column is None:
        return "A1"
    return f"{get_column_letter(column.source_columns[0])}{row}"
