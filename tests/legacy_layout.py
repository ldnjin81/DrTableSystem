"""Test helper: turns workbooks written in the old layout (field name, type and scope in rows
1-3, enums in <enum> sheets) into schema files, so tests can describe a table in one place.

* Each table sheet's header becomes ``<Table>.schema.xlsx``.
* Each ``<enum>Name`` sheet becomes ``<Enum>.enum.xlsx`` in the enum folder,
  its rows becoming the enum values. Extra columns (other than Id, Value and Comment) become a
  ``<Enum>Info`` table: a schema keyed by the enum and a new data workbook holding its rows.

"""

from __future__ import annotations

import re
from pathlib import Path

from harness import (
    DATA_ROW,
    IDENTIFIER_RE,
    NAME_ROW,
    SCHEMA_SUFFIXES,
    SCOPE_ROW,
    Schema,
    find_files,
    render_xlsx,
    strip_sheet_comment,
    table_name_of,
)
from openpyxl import Workbook, load_workbook

ENUM_SHEET_RE = re.compile(r"^<enum>(?P<name>[A-Za-z][A-Za-z0-9_]*)$")


def extract_schemas(
    input_path: Path,
    schema_dir: Path,
    enum_dir: Path,
    overwrite: bool = False,
) -> list[Path]:
    """Writes schema files (and <Enum>Info data workbooks) for the old sheets under input_path.

    The first sheet of a table, in file and sheet order, defines its schema. Existing files are
    kept unless ``overwrite`` is set. Returns the files written.
    """
    written: list[Path] = []
    seen: set[str] = set()
    for path, _ in find_files(input_path):
        if path.name.lower().endswith(SCHEMA_SUFFIXES):
            continue
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                match = ENUM_SHEET_RE.fullmatch(strip_sheet_comment(sheet.title))
                name = match.group("name") if match else table_name_of(sheet.title)
                key = f"<enum>{name}" if match else name
                if not name or not IDENTIFIER_RE.fullmatch(name) or key in seen:
                    continue
                seen.add(key)
                header = _header(sheet)
                if not header:
                    continue
                if match:
                    written += _convert_enum(sheet, name, header, schema_dir, enum_dir, path.parent, overwrite)
                else:
                    rows = [(0, n, t, s, None) for _, n, t, s in header]
                    written += _write(Schema("", name, rows), schema_dir, overwrite)
        finally:
            workbook.close()
    return written


def _convert_enum(
    sheet: object,
    name: str,
    header: list[tuple[int, object, object, object]],
    schema_dir: Path,
    enum_dir: Path,
    data_dir: Path,
    overwrite: bool,
) -> list[Path]:
    key_column = next((index for index, _, type_, _ in header if str(type_).strip().lower().startswith("id<")), 0)
    by_name = {str(field_name).strip(): index for index, field_name, _, _ in header}
    value_column = by_name.get("Value")
    comment_column = by_name.get("Comment")
    extras = [
        (index, field_name, type_, scope) for index, field_name, type_, scope in header
        if index not in (key_column, value_column, comment_column)
        and str(scope).strip() != "#"
    ]
    data = [list(row) for row in sheet.iter_rows(min_row=DATA_ROW, values_only=True)]
    data = [row for row in data if key_column < len(row) and row[key_column] not in (None, "")]

    def cell(row: list[object], index: int | None) -> object:
        return row[index] if index is not None and index < len(row) else None

    values = [(0, cell(row, key_column), cell(row, value_column), None, cell(row, comment_column)) for row in data]
    written = _write(Schema("", name, values, is_enum=True), enum_dir, overwrite)
    if extras:
        info = f"{name}Info"
        fields = [(0, "Id", f"ID<E{name}>", "all", None)] + [(0, n, t, s, None) for _, n, t, s in extras]
        written += _write(Schema("", info, fields), schema_dir, overwrite)
        target = data_dir / f"{info}.xlsx"
        if overwrite or not target.exists():
            workbook = Workbook()
            info_sheet = workbook.active
            info_sheet.title = info
            info_sheet.append(["Id", *(n for _, n, _, _ in extras)])
            info_sheet.append([])
            info_sheet.append([])
            for row in data:
                info_sheet.append([cell(row, key_column), *(cell(row, index) for index, _, _, _ in extras)])
            workbook.save(target)
            written.append(target)
    return written


def _write(schema: Schema, folder: Path, overwrite: bool) -> list[Path]:
    kind = "enum" if schema.is_enum else "schema"
    target = folder / f"{schema.name}.{kind}.xlsx"
    if target.exists() and not overwrite:
        return []
    folder.mkdir(parents=True, exist_ok=True)
    target.write_bytes(render_xlsx(schema))
    return [target]


def _header(sheet: object) -> list[tuple[int, object, object, object]]:
    """(0-based column, name, type, scope) from the old rows 1-3, up to the first empty name."""
    rows = list(sheet.iter_rows(min_row=NAME_ROW, max_row=SCOPE_ROW, values_only=True))
    rows += [()] * (SCOPE_ROW - len(rows))
    names, types, scopes = (list(row) for row in rows[:3])
    fields = []
    for index, name in enumerate(names):
        if name in (None, ""):
            break
        fields.append((
            index,
            name,
            types[index] if index < len(types) else None,
            scopes[index] if index < len(scopes) else None,
        ))
    return fields
