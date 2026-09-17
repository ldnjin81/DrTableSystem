from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from datagen.cli import main


def add_enum(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "<enum>Element"
    sheet.append(["Id", "Value", "Comment"])
    sheet.append(["ID<FName>", "int32", "FString"])
    sheet.append(["B", "B", "#"])
    sheet.append(["Fire", 0, "불"])
    sheet.append(["Water", None, "물"])


def add_table(workbook: Workbook, title: str = "Effects") -> object:
    sheet = workbook.create_sheet(title)
    sheet.append(
        [
            "Id",
            "Name",
            "Element",
            "ClientOnly",
            "ServerOnly",
            "Memo",
            "Reward[1]",
            "Reward[0]",
        ]
    )
    sheet.append(
        [
            "ID<int32>",
            "SubKey<FName>",
            "SubKey<EElement>",
            "float",
            "int64",
            "설명",
            "int32",
            "int32",
        ]
    )
    sheet.append(["B", "B", "B", "C", "S", "#", "B", "B"])
    sheet.append([1001, "Burn", "Fire", 12.5, 99, "무시", 20, 10])
    sheet.append([1002, "Freeze", "Water", None, 100, "무시", None, 30])
    return sheet


def save_valid(path: Path) -> None:
    workbook = Workbook()
    add_enum(workbook)
    add_table(workbook)
    ignored = workbook.create_sheet("#메모")
    ignored["A1"] = "이 시트는 형식과 무관하게 무시된다"
    workbook.save(path)


def output_bytes(root: Path) -> dict[str, bytes]:
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_build_outputs_scope_array_and_determinism(tmp_path: Path) -> None:
    source = tmp_path / "Tables.xlsx"
    save_valid(source)
    cpp = tmp_path / "cpp"
    client = tmp_path / "client"
    server = tmp_path / "server"
    args = [
        "build",
        "--input",
        str(source),
        "--out-cpp",
        str(cpp),
        "--out-client",
        str(client),
        "--out-server",
        str(server),
    ]
    assert main(args) == 0
    first = output_bytes(tmp_path)
    assert {path.name for path in cpp.iterdir()} == {
        "EDtElement.h",
        "DtEffectsRow.h",
        "DtGeneratedTables.h",
    }
    header = (cpp / "DtEffectsRow.h").read_text(encoding="utf-8")
    enum_header = (cpp / "EDtElement.h").read_text(encoding="utf-8")
    assert "Fire = 0, // 불" in enum_header
    assert "int32 Id = 0;" in header
    assert "FName Name;" in header
    assert "EDtElement Element = EDtElement::Fire;" in header
    assert "float ClientOnly = 0.0f;" in header
    assert "int32 Reward[2] = {};" in header
    assert "ClientOnly" in header
    assert "ServerOnly" not in header
    client_data = json.loads((client / "Effects.json").read_text(encoding="utf-8"))
    server_data = json.loads((server / "Effects.json").read_text(encoding="utf-8"))
    assert client_data["rows"][0]["Reward"] == [10, 20]
    assert client_data["rows"][1]["Reward"] == [30, 0]
    assert "ClientOnly" in client_data["rows"][0]
    assert "ServerOnly" not in client_data["rows"][0]
    assert "ServerOnly" in server_data["rows"][0]
    assert "ClientOnly" not in server_data["rows"][0]
    assert "Memo" not in client_data["rows"][0]
    assert [item["field"] for item in client_data["sub_keys"]] == ["Name", "Element"]
    assert main(args) == 0
    assert output_bytes(tmp_path) == first


def test_cpp_array_property_is_not_exposed_to_blueprint(tmp_path: Path) -> None:
    source = tmp_path / "Tables.xlsx"
    save_valid(source)
    output = tmp_path / "output"

    assert _build(source, output) == 0

    header = (output / "cpp" / "DtEffectsRow.h").read_text(encoding="utf-8")
    scalar_property = (
        '    UPROPERTY(EditAnywhere, BlueprintReadOnly, Category = "Dt|Effects")\n'
        "    int32 Id = 0;"
    )
    array_property = (
        '    UPROPERTY(EditAnywhere, Category = "Dt|Effects")\n'
        "    int32 Reward[2] = {};"
    )
    assert scalar_property in header
    assert array_property in header


def test_enum_info_table_and_stamp(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "<enum>ItemType"
    sheet.append(["Id", "Value", "Comment", "DisplayName", "MaxStack"])
    sheet.append(["ID<FName>", "int32", "FString", "FString", "int32"])
    sheet.append(["B", "B", "#", "C", "S"])
    sheet.append(["Weapon", 0, "무기류", "무기", 1])
    source = tmp_path / "enum-info.xlsx"
    workbook.save(source)
    result = main(
        [
            "build",
            "--input",
            str(source),
            "--out-cpp",
            str(tmp_path / "cpp"),
            "--out-client",
            str(tmp_path / "client"),
            "--out-server",
            str(tmp_path / "server"),
            "--stamp",
            "2026-09-17T00:00:00Z",
        ]
    )
    assert result == 0
    row_header = (tmp_path / "cpp" / "DtItemTypeInfoRow.h").read_text(encoding="utf-8")
    assert "EDtItemType Id = EDtItemType::Weapon;" in row_header
    assert "Value" not in row_header
    assert "Comment" not in row_header
    client = json.loads((tmp_path / "client" / "ItemTypeInfo.json").read_text(encoding="utf-8"))
    server = json.loads((tmp_path / "server" / "ItemTypeInfo.json").read_text(encoding="utf-8"))
    assert client["rows"] == [{"Id": "Weapon", "DisplayName": "무기"}]
    assert server["rows"] == [{"Id": "Weapon", "MaxStack": 1}]
    manifest = json.loads((tmp_path / "client" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["generated_at"] == "2026-09-17T00:00:00Z"


def test_cpp_scalar_initializers(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Numbers"
    sheet.append(["Id", "Big", "Ratio", "Precise", "Enabled", "Label"])
    sheet.append(["ID<int32>", "int64", "float", "double", "bool", "FString"])
    sheet.append(["B", "B", "B", "B", "B", "B"])
    sheet.append([1, 2, 3.5, 4.5, True, "값"])
    source = tmp_path / "initializers.xlsx"
    workbook.save(source)

    assert _build(source, tmp_path) == 0
    header = (tmp_path / "cpp" / "DtNumbersRow.h").read_text(encoding="utf-8")
    assert "int32 Id = 0;" in header
    assert "int64 Big = 0;" in header
    assert "float Ratio = 0.0f;" in header
    assert "double Precise = 0.0;" in header
    assert "bool Enabled = false;" in header
    assert "FString Label;" in header


def test_manifest_has_no_stamp_by_default(tmp_path: Path) -> None:
    source = tmp_path / "Tables.xlsx"
    save_valid(source)
    assert _build(source, tmp_path) == 0
    manifest = json.loads((tmp_path / "client" / "manifest.json").read_text(encoding="utf-8"))
    assert "generated_at" not in manifest
    assert "generated_at_utc" not in manifest
    assert manifest["source_files"] == ["Tables.xlsx"]


@pytest.mark.parametrize(
    ("mutate", "location", "message"),
    [
        (lambda sheet: sheet.__setitem__("A2", "FName"), "<enum>Element!A2", "기본키"),
        (
            lambda sheet: sheet.__setitem__("B2", "ID<int32>"),
            "<enum>Element!A2",
            "기본키",
        ),
        (lambda sheet: sheet.__setitem__("A4", "Bad-Name"), "<enum>Element!A4", "열거자"),
        (lambda sheet: sheet.__setitem__("A5", "Fire"), "<enum>Element!A5", "중복"),
    ],
)
def test_enum_validation_errors(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    mutate: object,
    location: str,
    message: str,
) -> None:
    workbook = Workbook()
    add_enum(workbook)
    mutate(workbook.active)
    source = tmp_path / "bad-enum.xlsx"
    workbook.save(source)
    assert main(["check", "--input", str(source)]) == 1
    stderr = capsys.readouterr().err
    assert location in stderr
    assert message in stderr


def test_enum_info_name_collision(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    workbook = Workbook()
    add_enum(workbook)
    enum_sheet = workbook.active
    enum_sheet["D1"] = "Label"
    enum_sheet["D2"] = "FString"
    enum_sheet["D3"] = "B"
    enum_sheet["D4"] = "불"
    enum_sheet["D5"] = "물"
    table = workbook.create_sheet("ElementInfo")
    table.append(["Id"])
    table.append(["ID<int32>"])
    table.append(["B"])
    table.append([1])
    source = tmp_path / "collision.xlsx"
    workbook.save(source)
    assert main(["check", "--input", str(source)]) == 1
    stderr = capsys.readouterr().err
    assert "<enum>Element!A1" in stderr
    assert "ElementInfo" in stderr


def test_sub_keys_can_be_empty(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    sheet.append(["Id", "Name"])
    sheet.append(["ID<FName>", "FString"])
    sheet.append(["B", "B"])
    sheet.append(["Sword", "검"])
    source = tmp_path / "input.xlsx"
    workbook.save(source)
    assert _build(source, tmp_path) == 0
    payload = json.loads((tmp_path / "client" / "Items.json").read_text(encoding="utf-8"))
    assert payload["sub_keys"] == []


def test_primary_key_scope_must_be_b(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workbook = Workbook()
    add_enum(workbook)
    sheet = add_table(workbook)
    sheet["A3"] = "C"
    source = tmp_path / "bad-primary-scope.xlsx"
    workbook.save(source)

    assert main(["check", "--input", str(source)]) == 1
    stderr = capsys.readouterr().err
    assert "Effects!A3" in stderr
    assert "기본키 범위는 B" in stderr


def test_check_does_not_write_files(tmp_path: Path) -> None:
    source = tmp_path / "Tables.xlsx"
    save_valid(source)
    before = output_bytes(tmp_path)
    assert main(["check", "--input", str(source)]) == 0
    assert output_bytes(tmp_path) == before


@pytest.mark.parametrize(
    ("mutate", "location", "message"),
    [
        (lambda sheet: sheet.__setitem__("A2", "int32"), "Effects!A2", "기본키"),
        (lambda sheet: sheet.__setitem__("B2", "ID<FName>"), "Effects!A2", "기본키"),
        (lambda sheet: sheet.__setitem__("A5", 1001), "Effects!A5", "중복"),
        (lambda sheet: sheet.__setitem__("A5", None), "Effects!A5", "비어"),
        (lambda sheet: sheet.__setitem__("C2", "EMissing"), "Effects!C2", "정의되지 않은"),
        (lambda sheet: sheet.__setitem__("D4", "숫자 아님"), "Effects!D4", "변환"),
        (lambda sheet: sheet.__setitem__("B1", "Id"), "Effects!B1", "중복"),
    ],
)
def test_validation_errors(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    mutate: object,
    location: str,
    message: str,
) -> None:
    workbook = Workbook()
    add_enum(workbook)
    sheet = add_table(workbook)
    mutate(sheet)
    source = tmp_path / "bad.xlsx"
    workbook.save(source)
    assert main(["check", "--input", str(source)]) == 1
    stderr = capsys.readouterr().err
    assert location in stderr
    assert message in stderr


@pytest.mark.parametrize(
    ("mutate", "location", "message"),
    [
        (lambda sheet: sheet.__setitem__("G1", "Reward[2]"), "Effects!G1", "연속"),
        (lambda sheet: sheet.__setitem__("G1", "Reward[0]"), "Effects!H1", "중복"),
        (lambda sheet: sheet.__setitem__("G2", "float"), "Effects!G2", "자료형"),
        (lambda sheet: sheet.__setitem__("G3", "C"), "Effects!G3", "범위"),
    ],
)
def test_array_validation_errors(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    mutate: object,
    location: str,
    message: str,
) -> None:
    workbook = Workbook()
    add_enum(workbook)
    sheet = add_table(workbook)
    mutate(sheet)
    source = tmp_path / "bad-array.xlsx"
    workbook.save(source)
    assert main(["check", "--input", str(source)]) == 1
    stderr = capsys.readouterr().err
    assert location in stderr
    assert message in stderr


def _build(source: Path, root: Path) -> int:
    return main(
        [
            "build",
            "--input",
            str(source),
            "--out-cpp",
            str(root / "cpp"),
            "--out-client",
            str(root / "client"),
            "--out-server",
            str(root / "server"),
        ]
    )
