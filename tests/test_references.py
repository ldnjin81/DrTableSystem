"""테이블 참조와 참조 검사 명령의 종단 테스트."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from harness import main
from openpyxl import Workbook


def _sheet(workbook: Workbook, name: str, headers: list[str], types: list[str],
           scopes: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    for record in (headers, types, scopes, *rows):
        sheet.append(record)


def _source(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Quests", ["Id", "Item", "Next[0]", "Next[1]", "ServerItem"],
           ["ID<int32>", "SubKey<Ref<Items>>", "Ref<Quests>", "Ref<Quests>",
            "Ref<Items>"], ["all", "all", "client", "client", "server"],
           [[1, 1001, 2, None, 1001], [2, None, None, None, None]])
    _sheet(workbook, "Items", ["Id", "Quest"], ["ID<int32>", "Ref<Quests>"],
           ["all", "all"], [[1001, 1]])
    _sheet(workbook, "Names", ["Id"], ["ID<name>"], ["all"], [["Sword"]])
    _sheet(workbook, "Links", ["Id", "NameRef"], ["ID<int64>", "ref<Names>"],
           ["all", "all"], [[3, None]])
    workbook.save(path)


def _build(source: Path, root: Path) -> int:
    return main(["build", "--input", str(source), "--out-cpp", str(root / "cpp"),
                 "--out-client", str(root / "client"), "--out-server", str(root / "server")])


def _snapshot(root: Path) -> dict[str, bytes]:
    return {str(path.relative_to(root)): path.read_bytes()
            for path in root.rglob("*") if path.is_file() and path.suffix != ".xlsx"}


def test_ref_build_graph_check_and_determinism(tmp_path: Path) -> None:
    source = tmp_path / "ref.xlsx"
    _source(source)
    assert _build(source, tmp_path) == 0
    client = json.loads((tmp_path / "client" / "manifest.json").read_text(encoding="utf-8"))
    server = json.loads((tmp_path / "server" / "manifest.json").read_text(encoding="utf-8"))
    assert [(r["table"], r["field"], r["target"], r["key_type"], r["array_length"],
             r["subkey"]) for r in client["references"]] == [
        ("Items", "Quest", "Quests", "int32", 1, False),
        ("Links", "NameRef", "Names", "name", 1, False),
        ("Quests", "Item", "Items", "int32", 1, True),
        ("Quests", "Next", "Quests", "int32", 2, False),
    ]
    assert [r["field"] for r in server["references"] if r["table"] == "Quests"] == [
        "Item", "ServerItem"
    ]
    quest = json.loads((tmp_path / "client" / "Quests.json").read_text(encoding="utf-8"))
    assert quest["rows"][1]["Next"] == [0, 0]
    assert quest["rows"][1]["Item"] == 0
    links = json.loads((tmp_path / "client" / "Links.json").read_text(encoding="utf-8"))
    assert links["rows"][0]["NameRef"] == ""
    link_header = (tmp_path / "cpp" / "DrLinksRow.h").read_text(encoding="utf-8")
    assert "FName NameRef = NAME_None;" in link_header
    header = (tmp_path / "cpp" / "DrQuestsRow.h").read_text(encoding="utf-8")
    assert 'meta = (TableRef = "Items")' in header
    assert 'meta = (TableRef = "Quests")' in header
    assert "int32 Item" in header
    assert "int32 Next[2]" in header
    assert main(["check", "--client", str(tmp_path / "client"),
                 "--server", str(tmp_path / "server")]) == 0
    graph = tmp_path / "graph.md"
    assert main(["graph", "--input", str(source), "--out", str(graph)]) == 0
    diagram = graph.read_text(encoding="utf-8")
    assert "flowchart LR" in diagram
    assert 'Names["Names (name)"]' in diagram
    assert 'Quests -->|"Next[2]"| Quests' in diagram
    assert 'Quests -->|"Item (SubKey)"| Items' in diagram
    first = _snapshot(tmp_path)
    assert _build(source, tmp_path) == 0
    assert main(["graph", "--input", str(source), "--out", str(graph)]) == 0
    assert _snapshot(tmp_path) == first


def test_check_reports_all_broken_refs_and_input_errors(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "ref.xlsx"
    _source(source)
    assert _build(source, tmp_path) == 0
    path = tmp_path / "client" / "Quests.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    data["rows"][0]["Item"] = 999
    data["rows"][0]["Next"] = [77, 88]
    path.write_text(json.dumps(data), encoding="utf-8")
    assert main(["check", "--client", str(tmp_path / "client")]) == 1
    output = capsys.readouterr().err
    assert "Quests.Item[1] = 999 → Items 테이블에 없음" in output
    assert "Quests.Next[1](0) = 77 → Quests 테이블에 없음" in output
    assert "Quests.Next[1](1) = 88 → Quests 테이블에 없음" in output
    assert main(["check", "--client", str(tmp_path / "missing")]) == 2


@pytest.mark.parametrize(("decl", "sheet", "cell"), [
    ("Ref<Missing>", "Quests", "B3"),
    ("Ref<<enum>Kind>", "Quests", "B3"),
    ("Ref<#Notes>", "Quests", "B3"),
    ("ID<Ref<Items>>", "Quests", "B2"),
    ("Ref<Items>=1001", "Quests", "B3"),
])
def test_ref_schema_errors(tmp_path: Path, capsys: pytest.CaptureFixture[str],
                           decl: str, sheet: str, cell: str) -> None:
    workbook = Workbook()
    workbook.active.title = "Quests"
    q = workbook.active
    q.append(["Id", "Item"])
    q.append([decl, "Ref<Items>"] if decl.startswith("ID<") else ["ID<int32>", decl])
    q.append(["all", "all"])
    q.append([1, 1001])
    _sheet(workbook, "Items", ["Id"], ["ID<int32>"], ["all"], [[1001]])
    workbook.create_sheet("#Notes")
    source = tmp_path / "bad.xlsx"
    workbook.save(source)
    assert _build(source, tmp_path) == 1
    assert f"[{sheet}.schema.xlsx]{sheet}!{cell}" in capsys.readouterr().err


def test_enum_ref_empty_is_error_and_info_is_target(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = Workbook()
    enum = workbook.active
    enum.title = "<enum>Kind"
    enum.append(["Id", "Value", "Label"])
    enum.append(["ID<name>", "int32", "string"])
    enum.append(["all", "all", "all"])
    enum.append(["Sword", 0, "검"])
    _sheet(workbook, "Uses", ["Id", "Kind", "Info"],
           ["ID<int32>", "Ref<KindInfo>", "Ref<KindInfo>"],
           ["all", "all", "all"], [[1, "Sword", None]])
    source = tmp_path / "enum.xlsx"
    workbook.save(source)
    assert _build(source, tmp_path) == 1
    assert "Uses!C4" in capsys.readouterr().err
    workbook["Uses"]["C4"] = "Sword"
    workbook.save(source)
    assert _build(source, tmp_path) == 0
    header = (tmp_path / "cpp" / "DrUsesRow.h").read_text(encoding="utf-8")
    assert "EDrKind Kind" in header
    assert 'TableRef = "KindInfo"' in header


def test_zero_key_warning(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "ref.xlsx"
    _source(source)
    # 기존 생성물의 키 집합에 0을 추가해 검사기의 충돌 경고를 확인한다.
    assert _build(source, tmp_path) == 0
    path = tmp_path / "client" / "Items.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    data["rows"].append({"Id": 0, "Quest": 0})
    path.write_text(json.dumps(data), encoding="utf-8")
    assert main(["check", "--client", str(tmp_path / "client")]) == 0
    assert "없음 값과 충돌" in capsys.readouterr().err



def _subkey_source(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "<enum>Kind", ["Id", "Value"], ["ID<name>", "int32"],
           ["all", "all"], [["A", 0], ["B", 1]])
    _sheet(workbook, "Monsters",
           ["Id", "DropGroup", "Groups[0]", "Groups[1]", "ByGroup", "Kind", "NameGroup"],
           ["ID<name>", "Ref<DropTable.GroupId>", "Ref<DropTable.GroupId>",
            "Ref<DropTable.GroupId>", "SubKey<Ref<DropTable.GroupId>>",
            "Ref<DropTable.KindKey>", "Ref<DropTable.NameKey>"],
           ["all", "all", "all", "all", "all", "all", "client"],
           [["Wolf", 10, 10, None, 10, "A", "Common"]])
    _sheet(workbook, "DropTable", ["Id", "GroupId", "NameKey", "KindKey", "Description"],
           ["ID<int32>", "SubKey<int32>", "SubKey<name>", "SubKey<EKind>", "string"],
           ["all", "all", "client", "all", "all"],
           [[1, 10, "Common", "A", "첫째"], [2, 10, "Common", "A", "둘째"]])
    workbook.save(path)


def test_subkey_ref_outputs_and_check(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "subkey.xlsx"
    _subkey_source(source)
    assert _build(source, tmp_path) == 0
    refs = json.loads((tmp_path / "client" / "manifest.json").read_text(encoding="utf-8"))[
        "references"
    ]
    assert {ref["field"]: (ref["target_key"], ref["cardinality"], ref["key_type"])
            for ref in refs} == {
        "ByGroup": ("GroupId", "many", "int32"),
        "DropGroup": ("GroupId", "many", "int32"),
        "Groups": ("GroupId", "many", "int32"),
        "Kind": ("KindKey", "many", "EKind"),
        "NameGroup": ("NameKey", "many", "name"),
    }
    server_refs = json.loads(
        (tmp_path / "server" / "manifest.json").read_text(encoding="utf-8")
    )["references"]
    assert "NameGroup" not in {ref["field"] for ref in server_refs}
    header = (tmp_path / "cpp" / "DrMonstersRow.h").read_text(encoding="utf-8")
    assert 'meta = (TableRef = "DropTable", TableRefKey = "GroupId")' in header
    assert 'meta = (TableRef = "DropTable", TableRefKey = "KindKey")' in header
    assert "EDrKind Kind" in header
    assert main(["check", "--client", str(tmp_path / "client"),
                 "--server", str(tmp_path / "server")]) == 0
    diagram = tmp_path / "graph.md"
    assert main(["graph", "--input", str(source), "--out", str(diagram)]) == 0
    assert "GroupId 1:N" in diagram.read_text(encoding="utf-8")
    path = tmp_path / "client" / "Monsters.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["rows"][0]["Groups"][0] = 404
    path.write_text(json.dumps(payload), encoding="utf-8")
    assert main(["check", "--client", str(tmp_path / "client")]) == 1
    assert "DropTable.GroupId에 해당 값 없음" in capsys.readouterr().err


@pytest.mark.parametrize(("field", "decl", "scope", "message"), [
    ("DropGroup", "Ref<DropTable.Id>", "all", "Ref<DropTable>를 쓰세요"),
    ("DropGroup", "Ref<DropTable.Description>", "all", "SubKey로 선언하세요"),
    ("DropGroup", "Ref<DropTable.Missing>", "all", "없습니다"),
    ("DropGroup", "Ref<DropTable.NameKey>", "all", "범위"),
    ("DropGroup", "ID<Ref<DropTable.GroupId>>", "all", "기본키"),
    ("DropGroup", "Ref<DropTable.GroupId>=10", "all", "기본값"),
])
def test_subkey_ref_schema_errors(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
    field: str, decl: str, scope: str, message: str
) -> None:
    source = tmp_path / "subkey.xlsx"
    _subkey_source(source)
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    sheet = workbook["Monsters"]
    sheet["B2"] = decl
    sheet["B3"] = scope
    workbook.save(source)
    assert _build(source, tmp_path) == 1
    output = capsys.readouterr().err
    assert "[Monsters.schema.xlsx]Monsters!B3" in output
    assert message in output


def test_subkey_enum_empty_and_warning(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "subkey.xlsx"
    _subkey_source(source)
    from openpyxl import load_workbook

    workbook = load_workbook(source)
    workbook["Monsters"]["F4"] = None
    workbook.save(source)
    assert _build(source, tmp_path) == 1
    assert "Monsters!F4" in capsys.readouterr().err
    workbook["Monsters"]["F4"] = "A"
    workbook["DropTable"].append([3, 0, "", "B", "없음 키"])
    workbook.save(source)
    assert _build(source, tmp_path) == 0
    assert main(["check", "--client", str(tmp_path / "client")]) == 0
    output = capsys.readouterr().err
    assert "DropTable.GroupId" in output
    assert "DropTable.NameKey" in output
    assert "없음 값과 충돌" in output



def test_chained_subkey_ref_resolves_final_type(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Uses", ["Id", "Value"], ["ID<int32>", "Ref<Top.Key>"],
           ["all", "all"], [[1, 7]])
    _sheet(workbook, "Top", ["Id", "Key"], ["ID<int32>", "SubKey<Ref<Middle.Key>>"],
           ["all", "all"], [[1, 7]])
    _sheet(workbook, "Middle", ["Id", "Key"], ["ID<int32>", "SubKey<Ref<Base>>"],
           ["all", "all"], [[1, 7]])
    _sheet(workbook, "Base", ["Id"], ["ID<int64>"], ["all"], [[7]])
    source = tmp_path / "chain.xlsx"
    workbook.save(source)
    assert _build(source, tmp_path) == 0
    header = (tmp_path / "cpp" / "DrUsesRow.h").read_text(encoding="utf-8")
    assert "int64 Value" in header
    manifest = json.loads((tmp_path / "client" / "manifest.json").read_text(encoding="utf-8"))
    ref = next(r for r in manifest["references"] if r["table"] == "Uses")
    assert (ref["target"], ref["target_key"], ref["key_type"]) == ("Top", "Key", "int64")
    assert main(["check", "--client", str(tmp_path / "client")]) == 0


def test_chained_subkey_type_cycle_reports_path(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "First", ["Id", "Key"], ["ID<int32>", "SubKey<Ref<Second.Key>>"],
           ["all", "all"], [[1, 1]])
    _sheet(workbook, "Second", ["Id", "Key"], ["ID<int32>", "SubKey<Ref<First.Key>>"],
           ["all", "all"], [[1, 1]])
    source = tmp_path / "cycle.xlsx"
    workbook.save(source)
    assert _build(source, tmp_path) == 1
    output = capsys.readouterr().err
    assert "자료형 순환: Second.Key → First.Key → Second.Key" in output
    assert "[First.schema.xlsx]First!B3" in output

